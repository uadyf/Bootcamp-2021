import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference

#Pandas Manipulation
data_luas = pd.read_csv('https://blog.sanbercode.com/wp-content/uploads/2020/08/luas-wilayah-menurut-kecamatan-di-kota-bandung-2017.csv')
data_jumlah = pd.read_csv('https://blog.sanbercode.com/wp-content/uploads/2020/08/jumlah-penduduk-kota-bandung.csv').sort_values(by='Kecamatan  ',ignore_index=True)

df = pd.DataFrame()
df['Nama Kecamatan'] = data_luas['Nama Kecamatan']
df['Kepadatan Penduduk'] = data_jumlah['Jumlah_Penduduk']/(data_luas['Luas Wilayah (m2)']/100)
df.to_excel('from_pandas.xlsx',index=False)

#Openpyxl Visualization
work = load_workbook('from_pandas.xlsx')
sheet = work.active

chart = BarChart()
chart.type = 'col'
chart.style = 1
chart.title = 'Kepadatan Penduduk'
chart.x_axis.title = 'Kecamatan'
chart.y_axis.title = 'Kepadatan per 100m2'

dt = Reference(sheet,min_col=2,max_col=2,min_row=2,max_row=31)
cat = Reference(sheet, min_col=1, max_col=1, min_row=2, max_row=31)
chart.height = 10
chart.width = 50
chart.add_data(dt)
chart.set_categories(cat)
sheet.add_chart(chart, "E2")

work.save('AuladiFachryAssidiqi.xlsx')